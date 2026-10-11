from pathlib import Path
import json,csv,collections,html,datetime
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
P=Path(__file__).resolve().parent
catalog=json.loads((P/'catalog.json').read_text()); corpus=json.loads((P/'corpus.json').read_text())
apps=[]
for scope in corpus['apps']:
 f=P/'apps'/(scope['slug']+'.json')
 if not f.exists():raise SystemExit('Missing '+str(f))
 d=json.loads(f.read_text());d['slug']=scope['slug'];d['kind']=scope.get('kind','app');d['cells_by_id']={c['component_id']:c for c in d['cells']};apps.append(d)
validation=json.loads((P/'evidence/validation.json').read_text())
if validation['errors']:raise SystemExit('Validation errors')
now=datetime.datetime.now(datetime.timezone.utc).isoformat(); payload={'created':now,'catalog':catalog,'apps':[{k:v for k,v in a.items() if k!='cells_by_id'} for a in apps],'validation':validation}
(P/'matrix.json').write_text(json.dumps(payload,indent=2)+'\n')
with (P/'matrix.csv').open('w',newline='') as f:
 w=csv.writer(f);w.writerow(['Ecosystem','Component']+[a['app'] for a in apps])
 for r in catalog:w.writerow([r['ecosystem'],r['name']]+[a['cells_by_id'][r['id']]['status'] for a in apps])
colors={'Y':'DCEFE4','N':'F9DEDC','M':'FFF0C4','NA':'ECEEF1'};ink={'Y':'166044','N':'972C2A','M':'805800','NA':'626B79'}
wb=Workbook();guide=wb.active;guide.title='Read me'
notes=[['RightKit cutover inventory','Source snapshot: '+now],['Y','Applicable capability routes through RightKit; no active competing local mechanism found within owner scope.'],['N','Applicable capability remains locally implemented.'],['M','RightKit & overlapping local implementation both remain active.'],['NA','Capability irrelevant to current app feature/runtime; reason appears in evidence.'],['How to use','Matrix rows are packages/crates/Swift products; columns are apps/tools. Click status to open evidence row; hover for rationale.'],['Migration work','Filtered list of N/M cells, concrete remaining work & implementation evidence.'],['Delivery flags','Version/path/publishing concerns stay separate from code cutover status.'],['Qualification','Static current-source audit; no builds, installed-binary tests or publication checks performed.'],['Snapshot','Concurrent checkouts can change. Evidence records hashes & excerpts; source drift is listed in validation.json.'],['Scope','Active apps/tools plus explicitly labelled archive/retained columns. Source roots & exclusions in Scope sheet / SCOPE.md.'],['Evidence','Each classification is a reviewer judgment grounded in cited source, not a dependency-count metric.'],['Editing','Workbook is completed audit snapshot. Editing a status does not change supporting evidence; re-audit cited callers before updating.']]
for x in notes:guide.append(x)
mx=wb.create_sheet('Matrix');mx.append(['Ecosystem','Component']+[a['app'] for a in apps]);mx.freeze_panes='C2'
ev=wb.create_sheet('Evidence');ev.append(['App','Component','Status','Rationale','Shared route','Local route','Applicability','Owner API','Remaining work','Delivery flags','Confidence'])
gaps=wb.create_sheet('Migration work');gaps.append(['App','Component','Status','Required work','Local route','Shared route','Owner API','Confidence'])
delivery=wb.create_sheet('Delivery flags');delivery.append(['App','Component','Status','Flags'])
scope=wb.create_sheet('Scope');scope.append(['App','Kind','Root','HEAD at reviewer read','Targets','Scope note'])
availability=json.loads((P/'evidence/availability.json').read_text()); availability_by_id={r['component_id']:r for r in availability['rows']}
cat=wb.create_sheet('Components');cat.append(['Ecosystem','Component','Version at catalog freeze','Description','Manifest','Owner source','Registry recorded (local ledger)','Staged (unverified)','Private'])
def evidence_text(arr):return '\n'.join(f"{e['path']}:{e['line']} — {e.get('symbol','')} {e.get('detail','')}" for e in arr)
evrows={}
for a in apps:
 scope.append([a['app'],a['kind'],a['root'],a.get('head',''),', '.join(a['targets']),a['scope_note']])
 for r in catalog:
  c=a['cells_by_id'][r['id']];key=(a['app'],r['id']);fields=[evidence_text(c.get(k,[])) for k in ['shared_evidence','local_evidence','applicability_evidence','owner_evidence']]
  ev.append([a['app'],r['name'],c['status'],c['rationale'],*fields,c['remaining_work'],'\n'.join(c.get('delivery_flags',[])),c.get('confidence','')]);evrows[key]=ev.max_row
  if c['status'] in ['N','M']:gaps.append([a['app'],r['name'],c['status'],c['remaining_work'],fields[1],fields[0],fields[3],c.get('confidence','')])
  if c.get('delivery_flags'):delivery.append([a['app'],r['name'],c['status'],'\n'.join(c['delivery_flags'])])
for r in catalog:
 mx.append([r['ecosystem'],r['name']]+[a['cells_by_id'][r['id']]['status'] for a in apps])
 for i,a in enumerate(apps,3):
  c=a['cells_by_id'][r['id']];cell=mx.cell(mx.max_row,i);cell.fill=PatternFill('solid',fgColor=colors[c['status']]);cell.font=Font(name='Arial',bold=True,color=ink[c['status']]);cell.alignment=Alignment(horizontal='center',vertical='center');cell.hyperlink=f"#'Evidence'!A{evrows[(a['app'],r['id'])]}";cell.comment=Comment(c['rationale']+'\n\n'+evidence_text(c.get('shared_evidence',[])+c.get('local_evidence',[])+c.get('applicability_evidence',[]))+'\n\nRemaining work: '+c['remaining_work'],'Source audit')
 av=availability_by_id[r['id']]
 cat.append([r['ecosystem'],r['name'],r['version'],r['description'],r['manifest'],r['source_root'],str(av['registry_recorded'] or 'No record'),str(av['staged_unverified'] or 'No record'),str(av['private_at_catalog_freeze'])])
for ws in wb:
 ws.sheet_view.showGridLines=False
 for row in ws:
  for cell in row:
   if ws.title!='Matrix' or cell.row==1 or cell.column<3:cell.font=Font(name='Arial',size=10)
   if ws.title!='Matrix':cell.alignment=Alignment(vertical='top',wrap_text=True)
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='20344B');cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=11);cell.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=46
 if ws.title!='Read me':ws.auto_filter.ref=ws.dimensions;ws.freeze_panes=ws.freeze_panes or 'C2'
 for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=28 if col<3 else 58
 ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.print_title_rows='1:1'
mx.column_dimensions['A'].width=13;mx.column_dimensions['B'].width=34
for i in range(3,mx.max_column+1):mx.column_dimensions[get_column_letter(i)].width=16
for row in range(2,mx.max_row+1):mx.row_dimensions[row].height=23
for ws in [ev,gaps,delivery]:
 ws.column_dimensions['C'].width=10
 for row in range(2,ws.max_row+1):ws.row_dimensions[row].height=75
for row in range(2,guide.max_row+1):guide.row_dimensions[row].height=48
guide.column_dimensions['A'].width=26;guide.column_dimensions['B'].width=116
wb.save(P/'rightkit-cutover-matrix.xlsx')
check=load_workbook(P/'rightkit-cutover-matrix.xlsx',data_only=False)
assert check['Matrix'].max_row==len(catalog)+1 and check['Matrix'].max_column==len(apps)+2
assert not any(c.data_type=='f' for ws in check for row in ws for c in row), 'No calculated workbook metrics used'
for i,r in enumerate(catalog,2):
 for j,a in enumerate(apps,3):assert check['Matrix'].cell(i,j).value==a['cells_by_id'][r['id']]['status']
# Markdown provides complete, portable report independent of spreadsheet software.
lines=['# RightKit cutover inventory','',f'{len(catalog)} components × {len(apps)} apps/tools = {len(catalog)*len(apps)} classified cells. Fresh production-source audit; Y/N/M/NA carry the definitions below.','',f'[Interactive matrix]({P}/matrix.html) · [Excel workbook]({P}/rightkit-cutover-matrix.xlsx) · [CSV]({P}/matrix.csv) · [JSON evidence]({P}/matrix.json)','', '| Code | Meaning |','|---|---|','| Y | Applicable capability uses RightKit; no active competing local mechanism found within owner scope. |','| N | Applicable capability remains locally implemented. |','| M | Shared & overlapping local mechanisms both active. |','| NA | Irrelevant to current feature/runtime, with reason recorded. |','', 'Source status & delivery status are separate. A local dependency override or stale pin does not turn Y into M. Product-specific adapters, schemas & UX remain local where shared owner does not supply that behavior.','', '## App inventory','', '| App/tool | Y | N | M | NA | Evidence & handoff |','|---|---:|---:|---:|---:|---|']
for a in apps:
 counts=collections.Counter(c['status'] for c in a['cells']);lines.append('| '+a['app']+' | '+' | '.join(str(counts[s]) for s in ['Y','N','M','NA'])+f" | [Report]({P}/apps/{a['slug']}.md) · [Full cells]({P}/apps/{a['slug']}.json) · [Handoff]({P}/handoffs/{a['slug']}.md) |")
lines+=['','## Evidence & scope','',f'- [Validation receipt]({P}/VALIDATION.md)',f'- [Cross-app migration queue]({P}/MIGRATION-QUEUE.md)',f'- [Source scope & exclusions]({P}/SCOPE.md)',f'- [Frozen app corpus]({P}/corpus.json)',f'- [Structural validation & concurrent source drift]({P}/evidence/validation.json)',f'- [Evidence hashes & quoted source context]({P}/evidence/citation-snapshots.json)',f'- [Reconciliation decisions]({P}/RECONCILIATION.md)',f'- [Unsupported-host reuse opportunities]({P}/UNSUPPORTED-HOST-REUSE.md)','','This audit describes inspected source routing. It does not certify installed binaries, runtime performance, build success or package publication. Use each N/M task to migrate callers, preserve product policy, remove proven overlap & verify affected targets.']
(P/'README.md').write_text('\n'.join(lines)+'\n')
for eco in ['cargo','npm','swift']:
 lines=['# '+eco+' cutover matrix','','Y = shared · N = local · M = mixed · NA = irrelevant. Evidence & tasks are in app reports / workbook.','', '| Component | '+' | '.join(a['app'] for a in apps)+' |','|---|'+'---|'*len(apps)]
 for r in catalog:
  if r['ecosystem']==eco:lines.append('| '+r['name']+' | '+' | '.join(a['cells_by_id'][r['id']]['status'] for a in apps)+' |')
 (P/f'matrix-{eco}.md').write_text('\n'.join(lines)+'\n')
# Shared-owner queue helps coordinate app chats without assigning future features.
queue=[]
for r in catalog:
 affected=[(a,a['cells_by_id'][r['id']]) for a in apps if a['kind'] not in ['archive','retained-source'] and a['cells_by_id'][r['id']]['status'] in ['N','M']]
 if affected:queue.append((r,affected))
queue.sort(key=lambda x:(-sum(c['status']=='M' for a,c in x[1]),-len(x[1]),x[0]['name']))
lines=['# Cross-app migration queue','','Active source only; archives & retained snapshots remain in full matrix. Grouping helps coordinate shared-owner work. Counts describe source classifications, not implementation effort.','', '| Component | Mixed apps | Still-local apps |','|---|---|---|']
for r,affected in queue:
 def app_links(status):return ', '.join(f"[{a['app']}]({P}/handoffs/{a['slug']}.md)" for a,c in affected if c['status']==status) or '—'
 lines.append(f"| {r['name']} | {app_links('M')} | {app_links('N')} |")
lines+=['','Work from each linked handoff: it identifies matching owner APIs, current callers, product behavior to preserve & separate delivery flags.']
(P/'MIGRATION-QUEUE.md').write_text('\n'.join(lines)+'\n')
# Browser UI has no external dependencies, requests or data upload.
template=(P/'matrix-template.html').read_text();small={'created':now,'catalog':[{k:r[k] for k in ['id','name','ecosystem','description','version']} for r in catalog],'apps':[{k:v for k,v in a.items() if k!='cells_by_id'} for a in apps]}
(P/'matrix.html').write_text(template.replace('__DATA__',json.dumps(small).replace('<','\\u003c')))
print('Built',len(apps),'apps,',len(catalog),'rows,',ev.max_row-1,'evidence cells; workbook round-trip matched all matrix values.')

# Per-app handoffs retain evidence & do not send to any chat.
handoff=P/'handoffs';handoff.mkdir(exist_ok=True)
for a in apps:
 lines=['# '+a['app']+' — RightKit migration handoff','','Source root: '+a['root'],'', 'Use this audit snapshot to plan cutovers; recheck cited callers before edits because other chats may have changed source. Select verified published versions; keep source cutover, packaging & execution qualification separate.','',f"[Full app evidence]({P}/apps/{a['slug']}.json)",'']
 for r in catalog:
  c=a['cells_by_id'][r['id']]
  if c['status'] not in ['N','M']:continue
  lines += ['## '+r['name']+' — '+c['status'],'',c['rationale'],'',c['remaining_work'],'']
  for key,label in [('local_evidence','Local'),('shared_evidence','Shared'),('owner_evidence','Owner')]:
   for e in c.get(key,[]):lines.append(f"- {label}: [{e.get('symbol') or Path(e['path']).name}]({e['path']}:{e['line']}) — {e.get('detail','')}")
  if c.get('delivery_flags'):lines += ['', 'Delivery: '+'; '.join(c['delivery_flags'])]
  lines += ['']
 if not any(c['status'] in ['N','M'] for c in a['cells']):lines += ['No direct N/M cutover is proven for current compatible hosts. Consult full evidence for runtime exclusions & unsupported reuse opportunities.','']
 lines += ['Before marking a task complete: migrate matching callers, remove only proven duplicate mechanisms, preserve product-specific schemas/policy/UX, then run affected target checks. Update this inventory from fresh evidence.']
 (handoff/(a['slug']+'.md')).write_text('\n'.join(lines)+'\n')
